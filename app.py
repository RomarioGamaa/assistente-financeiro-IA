import io
from contextlib import asynccontextmanager
import pandas as pd
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

from ai_service import processar_mensagem_financeira
import db

@asynccontextmanager
async def lifespan(app: FastAPI):
    db.init_db()
    yield

app = FastAPI(title="Assistente Financeiro IA", version="1.0.0", lifespan=lifespan)

class MensagemEntrada(BaseModel):
    telefone: str
    mensagem: str

@app.get("/")
def home():
    return FileResponse("chat.html")

@app.get("/exportar/{formato}")
def exportar_relatorio(formato: str, telefone: str = "web-demo"):
    dados = db.listar_transacoes_exportacao(telefone)
    
    if not dados:
        raise HTTPException(status_code=404, detail="Nenhum lançamento encontrado para este usuário.")

    df = pd.DataFrame(dados)
    
    # Padroniza e renomeia colunas
    df = df[["id", "tipo", "valor", "categoria", "forma_pagamento", "data_registro", "hora_registro"]]
    df.columns = ["ID", "Tipo", "Valor (R$)", "Categoria", "Forma de Pagamento", "Data", "Hora"]

    buffer = io.BytesIO()

    # Exportação em CSV
    if formato.lower() == "csv":
        df.to_csv(buffer, index=False, sep=";", encoding="utf-8-sig")
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=relatorio_financeiro.csv"}
        )

    # Exportação em Excel com Dashboard e Gráficos
    elif formato.lower() in ["xlsx", "excel"]:
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            # 1. Aba de Lançamentos detalhados
            df.to_excel(writer, sheet_name="Lançamentos", index=False)
            
            workbook = writer.book
            ws_dados = writer.sheets["Lançamentos"]
            
            # 2. Cria a aba "Dashboard" na primeira posição
            ws_dash = workbook.create_sheet(title="Dashboard", index=0)
            ws_dash.views.sheetView[0].showGridLines = True

            # Métricas e cálculos consolidados
            hoje = pd.Timestamp.now().strftime("%Y-%m-%d")
            total_entradas = float(df[df["Tipo"].str.upper() == "ENTRADA"]["Valor (R$)"].sum())
            total_saidas = float(df[df["Tipo"].str.upper() == "SAIDA"]["Valor (R$)"].sum())
            saldo_geral = total_entradas - total_saidas

            entradas_hoje = float(df[(df["Tipo"].str.upper() == "ENTRADA") & (df["Data"] == hoje)]["Valor (R$)"].sum())
            saidas_hoje = float(df[(df["Tipo"].str.upper() == "SAIDA") & (df["Data"] == hoje)]["Valor (R$)"].sum())
            saldo_hoje = entradas_hoje - saidas_hoje

            # Configurações de Estilos do Dashboard
            fill_topo = PatternFill("solid", fgColor="0F172A")
            fill_card = PatternFill("solid", fgColor="F8FAFC")
            font_titulo = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
            font_label = Font(name="Calibri", size=9, bold=True, color="64748B")
            font_valor = Font(name="Calibri", size=13, bold=True, color="0F172A")
            border_card = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1")
            )

            # Banner Superior
            ws_dash.merge_cells("A1:D2")
            cell_titulo = ws_dash["A1"]
            cell_titulo.value = "📊 RELATÓRIO EXECUTIVO FINANCEIRO"
            cell_titulo.fill = fill_topo
            cell_titulo.font = font_titulo
            cell_titulo.alignment = Alignment(horizontal="center", vertical="center")

            # Cards de Métricas (KPIs)
            kpis = [
                ("TOTAL ENTRADAS", total_entradas, "A"),
                ("TOTAL SAÍDAS", total_saidas, "B"),
                ("SALDO GERAL", saldo_geral, "C"),
                ("SALDO HOJE", saldo_hoje, "D")
            ]

            for label, valor, col in kpis:
                c_label = ws_dash[f"{col}4"]
                c_label.value = label
                c_label.font = font_label
                c_label.alignment = Alignment(horizontal="center", vertical="center")
                c_label.fill = fill_card
                c_label.border = border_card

                c_val = ws_dash[f"{col}5"]
                c_val.value = valor
                c_val.font = font_valor
                c_val.number_format = 'R$ #,##0.00'
                c_val.alignment = Alignment(horizontal="center", vertical="center")
                c_val.fill = fill_card
                c_val.border = border_card

            # Tabela de Apoio para o Gráfico
            ws_dash["A8"].value = "Tipo"
            ws_dash["B8"].value = "Total (R$)"
            ws_dash["A8"].font = Font(name="Calibri", size=10, bold=True)
            ws_dash["B8"].font = Font(name="Calibri", size=10, bold=True)

            ws_dash["A9"].value = "Entradas"
            ws_dash["B9"].value = total_entradas
            ws_dash["B9"].number_format = 'R$ #,##0.00'

            ws_dash["A10"].value = "Saídas"
            ws_dash["B10"].value = total_saidas
            ws_dash["B10"].number_format = 'R$ #,##0.00'

            # Gráfico de Barras Nativo
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = "Comparativo: Entradas vs Saídas"
            chart.y_axis.title = "Valor (R$)"
            chart.x_axis.title = "Fluxo"
            chart.height = 10
            chart.width = 15

            dados_ref = Reference(ws_dash, min_col=2, min_row=8, max_row=10)
            cats_ref = Reference(ws_dash, min_col=1, min_row=9, max_row=10)
            chart.add_data(dados_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            chart.legend = None

            ws_dash.add_chart(chart, "A12")

            for col_letter in ["A", "B", "C", "D"]:
                ws_dash.column_dimensions[col_letter].width = 22

            # 3. Estilização da aba "Lançamentos"
            fill_header = PatternFill("solid", fgColor="0F172A")
            font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

            for col_idx in range(1, len(df.columns) + 1):
                c = ws_dados.cell(row=1, column=col_idx)
                c.fill = fill_header
                c.font = font_header
                c.alignment = Alignment(horizontal="center", vertical="center")

            for row_idx in range(2, len(df) + 2):
                # Formata valor como moeda BRL
                cell_v = ws_dados.cell(row=row_idx, column=3)
                cell_v.number_format = 'R$ #,##0.00'
                cell_v.alignment = Alignment(horizontal="right")

                # Alinhamentos
                ws_dados.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
                ws_dados.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center")
                ws_dados.cell(row=row_idx, column=6).alignment = Alignment(horizontal="center")
                ws_dados.cell(row=row_idx, column=7).alignment = Alignment(horizontal="center")

            # Ajuste dinâmico da largura das colunas
            for col in ws_dados.columns:
                max_len = max(len(str(c.value or '')) for c in col)
                col_let = get_column_letter(col[0].column)
                ws_dados.column_dimensions[col_let].width = max(max_len + 4, 12)

        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=relatorio_financeiro.xlsx"}
        )

    raise HTTPException(status_code=400, detail="Formato inválido. Use 'csv' ou 'xlsx'.")

@app.post("/webhook")
def webhook(payload: MensagemEntrada):
    try:
        dados = processar_mensagem_financeira(payload.mensagem)
        resposta_texto = ""

        if dados.intencao == "REGISTRAR" and dados.itens:
            linhas_confirmacao = []
            for item in dados.itens:
                db.registrar_transacao(
                    telefone=payload.telefone,
                    tipo=item.tipo,
                    valor=item.valor,
                    categoria=item.categoria or "geral",
                    forma_pagamento=item.forma_pagamento
                )
                simbolo = "💰" if item.tipo == "ENTRADA" else "💸"
                linhas_confirmacao.append(f"• {simbolo} R$ {item.valor:.2f} ({item.tipo.title()} - {item.categoria or 'Geral'})")
            
            resposta_texto = "✅ Lançamentos registrados:\n" + "\n".join(linhas_confirmacao)

        elif dados.intencao == "CONSULTAR":
            totais = db.consultar_saldo_hoje(payload.telefone)
            resposta_texto = (
                f"📊 Balanço de Hoje:\n"
                f"• Entradas: R$ {totais['entradas']:.2f}\n"
                f"• Saídas: R$ {totais['saidas']:.2f}\n"
                f"• Saldo do Dia: R$ {totais['saldo']:.2f}"
            )
        else:
            resposta_texto = "Não consegui entender a operação. Tente algo como 'Recebi 50 no pix' ou 'Quanto fechei hoje?'"

        return {
            "remetente": payload.telefone,
            "extração": dados.model_dump(),
            "mensagem_resposta": resposta_texto
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))